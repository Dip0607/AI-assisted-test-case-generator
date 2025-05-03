import requests
from bs4 import BeautifulSoup
from docx import Document
from PyPDF2 import PdfReader
from io import BytesIO 

def extract_webpage_data(url):
    # Fetch the webpage content
    response = requests.get(url)

    if response.status_code == 200:
        # Parse the HTML content
        soup = BeautifulSoup(response.content, 'html.parser')

        # Extract the main content
        article_content = soup.find_all('p')
        content_text = "\n".join([p.get_text() for p in article_content])

        # Extract related links
        related_links = []
        for a in soup.find_all('a', href=True):
            related_links.append(a['href'])

        return content_text, related_links
    else:
        print(f"Failed to fetch the webpage. Status code: {response.status_code}")
        return None, None


def extract_urls(file):
    file_extension = file.filename.split(".")[-1].lower()
    if file_extension == "pdf":
        return None
    elif file_extension == "docx":
        return extract_urls_docx(file)
    else:
        raise ValueError("Unsupported file type. Please provide a PDF or DOCX file.")

def extract_urls_docx(docx_path):
    doc = Document(docx_path)
    urls = []
    for rel in doc.part.rels.values():
        if "hyperlink" in rel.reltype:
            urls.append(rel.target_ref)
    return urls



def read_file_content(file):
    file_extension = file.filename.split(".")[-1].lower()
    if file_extension == "pdf":
        return read_pdf(file)
    elif file_extension == "docx":
        return read_docx(file)
    else:
        raise ValueError("Unsupported file type. Please provide a PDF or DOCX file.")

def read_pdf(file):
    pdf_content = file.read()
    reader = PdfReader(BytesIO(pdf_content))
    return "".join(page.extract_text() for page in reader.pages)

def read_docx(file):
    docx_content = file.read()
    doc = Document(BytesIO(docx_content))
    return "\n".join(paragraph.text for paragraph in doc.paragraphs)


